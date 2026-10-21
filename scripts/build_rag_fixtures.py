"""Sinh lại tests/fixtures/sample.pdf và sample.docx (file cố định, đã commit).

Chạy: uv run python scripts/build_rag_fixtures.py
"""

from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

OUT = Path(__file__).resolve().parent.parent / "tests" / "fixtures"
FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"

TITLE = "Ghi chú đọc sách Sapiens"
SECTIONS = [
    (
        "Tóm tắt",
        "Sapiens của Yuval Noah Harari kể lại lịch sử loài người qua ba cuộc cách mạng: "
        "nhận thức, nông nghiệp và khoa học.",
    ),
    (
        "Liên hệ",
        "Mọi góp ý xin gửi về lan.nguyen@example.com hoặc gọi 0912 345 678. "
        "Thẻ thanh toán thử nghiệm: 4111 1111 1111 1111.",
    ),
]


def build_pdf() -> None:
    pdfmetrics.registerFont(TTFont("DejaVu", FONT))
    c = canvas.Canvas(str(OUT / "sample.pdf"), pagesize=A4)
    c.setTitle(TITLE)
    y = 800
    c.setFont("DejaVu", 16)
    c.drawString(50, y, TITLE)
    for heading, body in SECTIONS:
        y -= 40
        c.setFont("DejaVu", 12)
        c.drawString(50, y, heading)
        y -= 20
        c.setFont("DejaVu", 10)
        words, line = body.split(), ""
        for w in words:
            if len(line) + len(w) > 80:
                c.drawString(50, y, line)
                y -= 14
                line = ""
            line = f"{line} {w}".strip()
        c.drawString(50, y, line)
    c.save()


def build_docx() -> None:
    doc = Document()
    doc.core_properties.title = TITLE
    doc.add_heading(TITLE, level=0)
    for heading, body in SECTIONS:
        doc.add_heading(heading, level=1)
        doc.add_paragraph(body)
    doc.save(str(OUT / "sample.docx"))


if __name__ == "__main__":
    build_pdf()
    build_docx()
